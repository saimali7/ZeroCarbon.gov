"""Excel emissions report following the sheet structure of the EAD facility MRV reporting template."""
from __future__ import annotations

import math
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from model import (
    CH4_EF_DIESEL_COMB, CH4_EF_GAS_COMB, DIESEL_EF, DIESEL_NCV, FLARE_CE, GWP_CH4, IPCC_NG_EF, MONTHS, OPERATOR,
    SM3_PER_BOE, VERIFIER, YEAR, Facility,
)

TEAL = "0F4C5C"
HEAD_FILL = PatternFill("solid", fgColor=TEAL)
SUB_FILL = PatternFill("solid", fgColor="DCE8EB")
SECTION_FILL = PatternFill("solid", fgColor="EEF3F4")
INPUT_FILL = PatternFill("solid", fgColor="FFF8E1")
TOTAL_FILL = PatternFill("solid", fgColor="E8F0E3")
THIN = Side(style="thin", color="B7C4C8")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


class Sheet:
    def __init__(self, wb: Workbook, name: str, widths: list[float]):
        self.ws = wb.create_sheet(name)
        self.widths = widths
        self.n = len(widths)
        self.row = 1
        self.ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            self.ws.column_dimensions[get_column_letter(i)].width = w

    def _span(self, c1: int, c2: int) -> float:
        return sum(self.widths[c1 - 1:c2]) * 1.05

    def _height(self, text: str, width: float, size: float = 10) -> float:
        lines = sum(max(1, math.ceil(len(part) * size / 10 / max(width, 1))) for part in str(text).split("\n"))
        return max(15.0, lines * 13.2 + 3)

    def _merge(self, c1: int, c2: int) -> None:
        if c2 > c1:
            self.ws.merge_cells(start_row=self.row, start_column=c1, end_row=self.row, end_column=c2)

    def band(self, text: str, fill: PatternFill, font: Font, height: float) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.font, cell.alignment = font, Alignment(vertical="center", indent=1, wrap_text=True)
        for col in range(1, self.n + 1):
            self.ws.cell(self.row, col).fill = fill
        self._merge(1, self.n)
        self.ws.row_dimensions[self.row].height = height
        self.row += 1

    def title(self, text: str, f: Facility) -> None:
        self.band(text, HEAD_FILL, Font(bold=True, size=14, color="FFFFFF"), 28)
        self.band(f"{OPERATOR['name']}  |  {f.name}  |  EAD Facility ID {f.ead_id}  |  Reporting year {YEAR}",
                  SUB_FILL, Font(size=9.5, color="1D3B45"), 18)
        self.row += 1

    def section(self, text: str) -> None:
        self.band(text, SECTION_FILL, Font(bold=True, size=11, color="0F4C5C"), 20)

    def para(self, text: str, italic: bool = False, color: str = "1D2327", fill: PatternFill | None = None) -> None:
        cell = self.ws.cell(self.row, 1, text)
        cell.font, cell.alignment = Font(italic=italic, size=10, color=color), WRAP
        if fill:
            for col in range(1, self.n + 1):
                self.ws.cell(self.row, col).fill = fill
        self._merge(1, self.n)
        self.ws.row_dimensions[self.row].height = self._height(text, self._span(1, self.n))
        self.row += 1

    def kv(self, pairs: list[tuple[str, object]], label_cols: int = 1) -> None:
        for label, value in pairs:
            a = self.ws.cell(self.row, 1, label)
            a.font, a.alignment, a.border = Font(bold=True, size=10, color="33434A"), WRAP, BORDER
            for col in range(2, label_cols + 1):
                self.ws.cell(self.row, col).border = BORDER
            b = self.ws.cell(self.row, label_cols + 1, value)
            b.fill, b.alignment, b.border = INPUT_FILL, WRAP, BORDER
            if isinstance(value, (int, float)):
                b.number_format = "#,##0" if isinstance(value, int) else "#,##0.0#"
                b.alignment = Alignment(horizontal="left", vertical="top")
            for col in range(label_cols + 1, self.n + 1):
                self.ws.cell(self.row, col).fill = INPUT_FILL
                self.ws.cell(self.row, col).border = BORDER
            if label_cols > 1:
                self.ws.merge_cells(start_row=self.row, start_column=1, end_row=self.row, end_column=label_cols)
            self._merge(label_cols + 1, self.n)
            h = max(self._height(label, self._span(1, label_cols)), self._height(str(value), self._span(label_cols + 1, self.n)))
            self.ws.row_dimensions[self.row].height = h
            self.row += 1

    def table(self, headers: list[str], rows: list[list], fmts: dict[int, str] | None = None,
              input_cols: set[int] | None = None, total_last: bool = False, spans: list[int] | None = None) -> tuple[int, int]:
        fmts = fmts or {}
        input_cols = input_cols if input_cols is not None else set(range(len(headers)))
        spans = spans or [1] * len(headers)
        starts = [1 + sum(spans[:i]) for i in range(len(spans))]

        def write_row(values: list, style) -> None:
            heights = []
            for i, value in enumerate(values):
                c1, c2 = starts[i], starts[i] + spans[i] - 1
                cell = self.ws.cell(self.row, c1, value)
                style(cell, i, value)
                for col in range(c1 + 1, c2 + 1):
                    extra = self.ws.cell(self.row, col)
                    extra.border, extra.fill = BORDER, copy(cell.fill)
                if c2 > c1:
                    self.ws.merge_cells(start_row=self.row, start_column=c1, end_row=self.row, end_column=c2)
                heights.append(self._height("" if value is None else str(value), self._span(c1, c2)))
            self.ws.row_dimensions[self.row].height = max(heights)
            self.row += 1

        def header_style(cell, i, value) -> None:
            cell.font, cell.fill, cell.border = Font(bold=True, size=9.5, color="FFFFFF"), HEAD_FILL, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="center")

        write_row(headers, header_style)
        first = self.row
        for r_i, row in enumerate(rows):
            is_total = total_last and r_i == len(rows) - 1

            def body_style(cell, i, value, is_total=is_total) -> None:
                cell.border, cell.alignment = BORDER, WRAP
                if isinstance(value, (int, float)) and not isinstance(value, bool):
                    cell.number_format = fmts.get(i, "#,##0")
                    cell.alignment = Alignment(horizontal="right", vertical="top")
                if is_total:
                    cell.fill, cell.font = TOTAL_FILL, Font(bold=True)
                elif i in input_cols:
                    cell.fill = INPUT_FILL

            write_row(row, body_style)
        return first, self.row - 1

    def gap(self, n: int = 1) -> None:
        self.row += n

    def dropdown(self, options: list[str], col: int, first: int, last: int) -> None:
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        self.ws.add_data_validation(dv)
        letter = get_column_letter(col)
        dv.add(f"{letter}{first}:{letter}{last}")


def build_report(f: Facility, out: Path) -> None:
    c = f.calc
    wb = Workbook()
    wb.remove(wb.active)
    wb.properties.creator = f"{OPERATOR['name']} ({f.ghg_lead})"
    wb.properties.title = f"EAD MRV Emissions Report RY{YEAR} - {f.short}"
    wb.properties.company = OPERATOR["name"]

    sheets_index = [
        ("A_Contents", "Table of contents and document status", "Completed"),
        ("B_Guidance", "Summary of reporting rules applied", "Completed"),
        ("C1_Identifiers", "Operator, facility and contact identifiers; operator declaration", "Completed"),
        ("C2_Facility_Description", "Facility description, technical units, static/dynamic data, approach selection, emissions summary, production", "Completed"),
        ("D1_Source_Streams", "Source streams (calculation-based), categorisation and tiers", "Completed"),
        ("D2_Calculation_Approach", "Activity data, calculation factors and CO2 per source stream; monthly activity data", "Completed"),
        ("E1_Emission_Sources_Measured", "Emission sources monitored by measurement (CEMS)", "Not applicable"),
        ("E2_Measurement_Approach", "Measurement-based methodology details", "Not applicable"),
        ("F_Fallback_Approach", "Fallback methodology and justification", "Not applicable"),
        ("G_Methane", "Methane emission sources and quantification", "Completed"),
        ("H1_Verification_Data_Gaps", "Data gaps, substitution and verification", "Completed"),
        ("I_Management_QA", "Roles, procedures and measuring-instrument QA/QC", "Completed"),
        ("J_Mitigation_Measures", "Emission reduction actions, plans and studies; offsets (not netted)", "Completed"),
        ("K_Reference_Lists", "Reference lists, default factors and GWP values", "Reference"),
    ]

    s = Sheet(wb, "A_Contents", [34, 70, 18])
    s.title("EAD Facility-Level MRV Reporting Template: Annual Emissions Report", f)
    s.kv([("Operator", OPERATOR["name"]), ("Facility", f.name), ("Reporting period", f"1 January {YEAR} to 31 December {YEAR}"),
          ("Monitoring Plan reference", f"{f.code}-HSE-MP-001 Rev {f.mp_rev} ({f.mp_date:%d %b %Y})"),
          ("Workbook version / date", f"Final for submission, {f.submitted:%d %b %Y}")])
    s.gap()
    s.section("Contents")
    s.table(["Sheet", "Content", "Status"], [list(r) for r in sheets_index], input_cols=set())
    s.gap()
    s.para("Workbook structure recreated for the ZeroCarbon.gov demo from EAD's Technical Guidance for MRV (Appendix 1: Manual for EAD Template "
           "Workbook) and the EAD facility-level MRV workshop of 12 March 2026. It is not the official EAD workbook. Yellow cells are operator inputs.",
           italic=True, color="5F6B72")

    s = Sheet(wb, "B_Guidance", [110])
    s.title("B. Guidance applied when completing this workbook", f)
    for line in [
        "Scope: direct (Scope 1) emissions of carbon dioxide (CO2) and methane (CH4) from the facility, under the operational control approach.",
        "Threshold and sectors: facilities in covered sectors (power, oil & gas, industry) emitting 25,000 t CO2e or more per year must register and report.",
        "Deadline: the verified annual emissions report is due by 31 March of the year after the reporting period (grace period to 14 April).",
        "Methodologies: calculation-based (activity data x emission factor x oxidation factor), mass balance, or measurement-based (CEMS). A fallback approach "
        "may be used only with justification and an overall uncertainty of no more than 7.5%.",
        "Tiers for activity data (maximum permissible uncertainty): Tier 1 +/-7.5%, Tier 2 +/-5.0%, Tier 3 +/-2.5%, Tier 4 +/-1.5%. Major source streams should meet at least Tier 2.",
        "Source stream categories: De-minimis = jointly <1 kt CO2e or <2% of total (up to 20 kt); Minor = jointly <5 kt CO2e or <10% of total (up to 100 kt); Major = all others.",
        "Completeness: monitoring and reporting must cover all emissions, avoid double counting and prevent data gaps. Any deviation from the Monitoring Plan and any data gap must be reported (sheet H1).",
        f"Global warming potential: CH4 = {GWP_CH4} (IPCC AR5, 100-year), consistent with UNFCCC reporting under the Paris Agreement.",
        "Offsets and carbon removals are reported for information only (sheet J) and are not netted against reported emissions.",
        "Standard conditions: gas volumes in this workbook are at 15 degC and 101.325 kPa (Sm3).",
    ]:
        s.para("- " + line)

    s = Sheet(wb, "C1_Identifiers", [44, 80])
    s.title("C1. Identifiers", f)
    s.section("(a) Operator")
    s.kv([("Operator legal name", OPERATOR["name"]), ("Operator name (Arabic)", OPERATOR["name_ar"]),
          ("Commercial licence number", OPERATOR["licence"]), ("Registered address", OPERATOR["address"]),
          ("Consolidation approach", "Operational control (Dunes Energy Company operates and controls 100% of the facility)")])
    s.gap()
    s.section("(b) Facility")
    s.kv([("Facility name", f.name), ("EAD facility registration ID", f.ead_id), ("EAD environmental permit number", f.permit),
          ("Covered sector / activity", "Oil & Gas: onshore crude oil production, separation, stabilisation, gas compression and dehydration"),
          ("Location", f"{f.field_name}, Al Dhafra Region, Emirate of Abu Dhabi"),
          ("Coordinates (WGS84, facility centre)", f"{f.lat:.4f} N, {f.lon:.4f} E"),
          ("Start of operations", str(f.start_year)), ("New facility (first 24 months of operation)?", "No"),
          ("Site schematic reference", f"Monitoring Plan {f.code}-HSE-MP-001 Rev {f.mp_rev}, section 3, Figure 1")])
    s.gap()
    s.section("(c) Contacts")
    s.kv([("GHG manager / person in charge", f"{f.ghg_lead}, GHG & Energy Lead"), ("Email", f.ghg_lead_email), ("Telephone", f.ghg_lead_phone),
          ("Facility manager", f.facility_manager), ("Corporate HSE Director", OPERATOR["hse_director"])])
    s.gap()
    s.section("(d) Reporting")
    s.kv([("Reporting period", f"01-Jan-{YEAR} to 31-Dec-{YEAR}"), ("Monitoring Plan applied", f"Rev {f.mp_rev}, submitted to EAD on {f.mp_date:%d %b %Y}"),
          ("Accredited verifier", f"{VERIFIER['name']} ({VERIFIER['ead_listing']})"), ("Date of submission", f"{f.submitted:%d %b %Y}")])
    s.gap()
    s.section("(e) Operator declaration")
    s.para("I declare that the information in this report is complete and accurate to the best of my knowledge, has been prepared in accordance with "
           "the approved Monitoring Plan and EAD's Technical Guidance for MRV, and that all emissions from the facility have been reported.")
    s.kv([("Name / position", f"{f.facility_manager}, Facility Manager"), ("Date", f"{f.submitted:%d %b %Y}")])

    s = Sheet(wb, "C2_Facility_Description", [16, 44, 22, 22, 18, 16])
    s.title("C2. Facility description and summary", f)
    s.section("(a) Non-technical summary")
    s.para(facility_summary(f))
    s.gap()
    s.section("(b) Main technical units")
    s.table(["Tag", "Description", "Capacity", "Source stream / fuel", "In scope", "Notes"], technical_units(f), input_cols={0, 1, 2, 3, 4, 5})
    s.gap()
    s.section("(c) Static data")
    s.kv([("Installed power generation", "3 x 24 MW gas turbine generators (GTG-A/B/C)"),
          ("Installed compression", "2 x 11 MW gas-turbine-driven centrifugal compressors (K-201A/B)"),
          ("Design crude processing capacity", "75,000 bbl/d" if f.flagged else "55,000 bbl/d"),
          ("Flare system", "HP flare FL-501 and LP flare FL-502, continuous pilots, ultrasonic flow meters FT-5101/FT-5102")], label_cols=2)
    s.gap()
    s.section("(d) Dynamic data (reporting year)")
    s.kv([("Crude oil produced (bbl)", c["oil_bbl"]), ("Average crude rate (bbl/d)", round(c["oil_bbl"] / 365)),
          ("Sales gas exported (Sm3)", c["gas_exported_sm3"]), ("Sales gas exported (boe, 5,800 scf/boe)", round(c["gas_boe"])),
          ("Total hydrocarbon production (MMboe)", c["mmboe"]), ("Fuel gas consumed (Sm3)", f.fuel_gas_sm3),
          ("Gas flared (Sm3)", c["flare_sm3"]), ("Diesel consumed (t)", c["diesel_t"]), ("Operating hours", 8_760)], label_cols=2)
    s.gap()
    s.section("(e) Monitoring approaches used")
    first, last = s.table(["Approach", "Used?", "Sheets"], [
        ["Calculation approach for CO2", "Yes", "D1, D2"], ["Measurement approach for CO2", "No", "E1, E2"],
        ["Fallback approach", "No", "F"], ["Methane emissions", "Yes", "G"]], input_cols={1})
    s.dropdown(["Yes", "No"], 2, first, last)
    s.gap()
    s.section("(f) Emissions summary")
    rows = [[sid, name, co2, 0.0, co2, co2 / c["total_co2e"]] for sid, name, co2, _ in c["streams"]]
    rows.append(["CH4", "All methane sources (sheet G)", 0, c["ch4_total"], c["ch4_co2e"], c["ch4_co2e"] / c["total_co2e"]])
    rows.append(["TOTAL", "Facility total (Scope 1)", c["co2_total"], c["ch4_total"], c["total_co2e"], 1.0])
    s.table(["ID", "Source", "CO2 (t)", "CH4 (t)", "Total (t CO2e)", "Share"], rows,
            fmts={2: "#,##0", 3: "#,##0.0", 4: "#,##0", 5: "0.0%"}, input_cols={2, 3, 4}, total_last=True)
    s.gap()
    s.section("(g) Product / production information (collected from RY2025 for sector benchmarking)")
    first, last = s.table(["Category", "Product", "Quantity", "Unit", "Emissions attributed", "Share"], [
        ["Other", "Stabilised crude oil", c["oil_bbl"], "bbl", "Allocated by energy", "88%"],
        ["Other", "Sales gas", c["gas_exported_sm3"], "Sm3", "Allocated by energy", "12%"]], input_cols={0, 1, 2, 3, 4, 5})
    s.dropdown(["Refinery products", "Hydrogen", "Synthesis gas", "Heat benchmark", "Fuel benchmark", "Process emissions", "Other"], 1, first, last)
    s.gap()
    s.section("(h) Offsets and removals")
    s.para("None. Offsets and carbon removals, if any, are reported in sheet J for information only and are not netted against the totals above.")

    s = Sheet(wb, "D1_Source_Streams", [8, 18, 13, 34, 14, 13, 28, 11, 12, 34])
    s.title("D1. Source streams (calculation-based methodology)", f)
    rows = []
    for sid, name, co2, cat in c["streams"]:
        meta = STREAM_META[sid]
        rows.append([sid, name, meta["type"], meta["sources"], co2, cat, meta["instrument"], meta["tier"], meta["unc"], meta["factors"]])
    first, last = s.table(["ID", "Source stream", "Type", "Emission sources", "Est. emissions (t CO2)", "Category",
                           "Activity data measurement", "Tier (activity data)", "Max. uncertainty", "Calculation factors basis"], rows)
    s.dropdown(["Major", "Minor", "De-minimis"], 6, first, last)
    s.dropdown(["Tier 1", "Tier 2", "Tier 3", "Tier 4", "No tier (fallback)"], 8, first, last)
    s.gap()
    s.para("Categorisation per EAD Technical Guidance (Appendix 1.3, Step 3). Emission source and meter tags refer to Monitoring Plan Figure 1.",
           italic=True, color="5F6B72")

    s = Sheet(wb, "D2_Calculation_Approach", [8, 16, 15, 9, 11, 11, 12, 11, 13, 12, 13, 16, 30])
    s.title("D2. Calculation approach: activity data, calculation factors and emissions", f)
    s.section("(a) Annual calculation per source stream")
    ce_note = f"CE {FLARE_CE:.2f} in EF"
    rows = [
        ["SS-01", "Fuel gas", f.fuel_gas_sm3, "Sm3", c["fuel_ncv"], "MJ/Sm3", c["fuel_tj"], c["fuel_ef_used"], "t CO2/TJ", 1.0, c["fuel_co2"],
         "AD Tier 3 / NCV Tier 3 / EF Tier 3", "FT-3001; NCV and EF site-specific from quarterly GC analysis"],
        ["SS-02", "HP flare gas", c["hp_sm3"], "Sm3", "n/a", "", "n/a", c["flare_co2_f"], "t CO2/10^3 Sm3", ce_note, c["hp_co2"],
         "AD Tier 2 / EF Tier 3", "FT-5101; EF from quarterly GC of flare header gas"],
        ["SS-03", "LP flare gas", c["lp_sm3"], "Sm3", "n/a", "", "n/a", c["flare_co2_f"], "t CO2/10^3 Sm3", ce_note, c["lp_co2"],
         "AD Tier 2 / EF Tier 3", "FT-5102; EF as SS-02"],
        ["SS-04", "Diesel", c["diesel_t"], "t", DIESEL_NCV, "GJ/t", c["diesel_tj"], DIESEL_EF, "t CO2/TJ", 1.0, c["diesel_co2"],
         "AD Tier 1 / NCV, EF IPCC default", "Supplier invoices + stock change (see (c))"],
        ["", "TOTAL CO2", "", "", "", "", "", "", "", "", c["co2_total"], "", ""],
    ]
    s.table(["ID", "Source stream", "Activity data", "Unit", "NCV", "NCV unit", "Energy (TJ)", "Emission factor", "EF unit",
             "Oxidation / combustion factor", "CO2 (t)", "Tiers", "Data source"], rows,
            fmts={2: "#,##0.0", 4: "0.00", 6: "#,##0.0", 7: "0.000", 9: "0.00", 10: "#,##0"}, total_last=True)
    s.para("Combustion: CO2 = activity data (Sm3) x NCV (MJ/Sm3) x 10^-6 x EF (t CO2/TJ) x oxidation factor.  "
           "Flaring: CO2 = volume (Sm3) / 1000 x EF (t CO2/10^3 Sm3), where EF = [sum(x_i n_i) x CE + x_CO2] x 44.01 / 23.645.  "
           "Diesel: CO2 = mass (t) x NCV (GJ/t) / 1000 x EF (t CO2/TJ).", italic=True, color="5F6B72")
    s.gap()
    s.section("(b) Monthly activity data")
    rows = [[MONTHS[m], c["fuel_monthly"][m], c["hp_m"][m], c["lp_m"][m], c["hp_m"][m] + c["lp_m"][m]] for m in range(12)]
    rows.append(["Total", f.fuel_gas_sm3, c["hp_sm3"], c["lp_sm3"], c["flare_sm3"]])
    s.table(["Month", "SS-01 Fuel gas (Sm3)", "SS-02 HP flare (Sm3)", "SS-03 LP flare (Sm3)", "Flare total (Sm3)"], rows,
            total_last=True, spans=[1, 2, 2, 2, 2])
    s.gap()
    s.section("(c) Diesel stock reconciliation")
    s.kv([("Opening stock 01-Jan (t)", f.diesel_open_t), ("Deliveries per invoices (t)", c["diesel_deliveries_t"]),
          ("Closing stock 31-Dec (t)", f.diesel_close_t), ("Consumption (t)", c["diesel_t"])], label_cols=3)

    for name, title in [("E1_Emission_Sources_Measured", "E1. Emission sources monitored by measurement"),
                        ("E2_Measurement_Approach", "E2. Measurement-based methodology")]:
        s = Sheet(wb, name, [110])
        s.title(title, f)
        s.para("Not applicable. No emission sources at the facility are monitored using a measurement-based methodology (no CEMS installed on covered sources).")

    s = Sheet(wb, "F_Fallback_Approach", [110])
    s.title("F. Fallback approach", f)
    s.para("Not applicable. All source streams are monitored using the calculation-based methodology at the tiers stated in sheet D1. "
           "No fallback methodology has been applied during the reporting period.")

    s = Sheet(wb, "G_Methane", [7, 32, 15, 34, 34, 11, 8, 12, 30])
    s.title("G. Methane (CH4) emissions", f)
    rows = [[m["id"], m["source"], m["category"], m["method"], m["basis"],
             m["ch4"] if m["ch4"] is not None else "n/a", GWP_CH4 if m["ch4"] is not None else "",
             round(m["ch4"] * GWP_CH4) if m["ch4"] is not None else "n/a", m["note"]] for m in c["methane"]]
    rows.append(["", "TOTAL METHANE", "", "", "", c["ch4_total"], GWP_CH4, c["ch4_co2e"], ""])
    s.table(["ID", "Methane source", "Category", "Quantification method", "Basis / activity", "CH4 (t)", "GWP", "t CO2e", "Notes"],
            rows, fmts={5: "#,##0.0", 7: "#,##0"}, total_last=True)
    s.para(f"GWP for CH4 = {GWP_CH4} (IPCC AR5). Combustion factors: IPCC 2006 Guidelines Vol. 2 Table 2.2 "
           f"({CH4_EF_GAS_COMB:g} kg CH4/TJ natural gas, {CH4_EF_DIESEL_COMB:g} kg CH4/TJ diesel).", italic=True, color="5F6B72")

    s = Sheet(wb, "H1_Verification_Data_Gaps", [22, 20, 34, 34, 16, 12])
    s.title("H1. Verification and data gaps", f)
    s.section("(a) Data gaps during the reporting period")
    if f.flagged:
        s.table(["Source stream", "Period", "Cause", "Substitution method", "Est. impact (t CO2e)", "% of total"],
                [["None", "-", "No data gaps identified during the reporting period.", "-", 0, 0.0]], fmts={5: "0.00%"})
        s.para("All source streams were monitored in accordance with Monitoring Plan Rev 3.0 throughout the reporting period.")
    else:
        gap_sm3 = sum(r["lp"] for r in c["log"] if r["lp_src"].startswith("SUBSTITUTED"))
        impact = round(gap_sm3 / 1000 * (c["flare_co2_f"] + c["flare_ch4_f"] * GWP_CH4))
        s.table(["Source stream", "Period", "Cause", "Substitution method", "Est. impact (t CO2e)", "% of total"],
                [["SS-03 LP flare gas", "11-13 Mar 2025 (3 days)", "FT-5102 transducer fault; meter offline until transducer replaced and recalibrated on 14-Mar-2025.",
                  "Average of the preceding 30 days of metered data (Monitoring Plan section 8.3). EAD notified 02-Apr-2025.", impact, impact / c["total_co2e"]]],
                fmts={5: "0.00%"})
    s.gap()
    s.section("(b) Verification")
    s.kv([("Verification body", VERIFIER["name"]), ("Accreditation", VERIFIER["accreditation"]),
          ("Verification statement reference", f"{f.verifier_ref}, dated {f.verif_date:%d %b %Y}"),
          ("Level of assurance", "Reasonable"),
          ("Verification outcome", "Verified (see enclosed statement)" if f.flagged else "Verified: unmodified opinion (see enclosed statement)"),
          ("Site visit", f"{f.site_visit:%d %b %Y}")], label_cols=2)
    s.gap()
    s.section("(c) Self-verification and internal quality assurance")
    s.table(["Source stream", "Control", "Description", "Evidence", "Frequency", "Owner"], [
        ["SS-01 Fuel gas", "Meter vs gas balance", "FT-3001 monthly total reconciled with production allocation fuel gas figure (tolerance 2%).", "Monthly GHG workbook", "Monthly", "GHG Lead"],
        ["SS-02/03 Flare", "Meter vs gas balance", "Flare meter totals reconciled with gas balance flare figure; deviations >10% investigated.", "Monthly GHG workbook", "Monthly", "GHG Lead"],
        ["SS-04 Diesel", "Invoice check", "Deliveries reconciled with supplier invoices and tank dips.", "Invoices, dip sheets", "Monthly", "Materials"],
        ["All", "Year-on-year review", "Annual emissions and intensities compared with previous years; variances >10% explained.", "Annual review memo", "Annual", "HSE Manager"],
    ], input_cols={2})

    s = Sheet(wb, "I_Management_QA", [20, 30, 22, 12, 14, 14, 16, 16])
    s.title("I. Management and quality assurance", f)
    s.section("(a) Roles and responsibilities")
    s.table(["Post", "Responsibility"], [
        ["Facility Manager", "Accountable for MRV compliance; signs operator declaration."],
        ["GHG & Energy Lead", "Prepares Monitoring Plan and emissions report; runs monthly reconciliations; liaises with EAD and verifier."],
        ["Instrumentation Supervisor", "Maintains and calibrates flow meters FT-3001, FT-5101, FT-5102 per calibration schedule."],
        ["Laboratory Coordinator", "Arranges quarterly gas sampling and third-party GC analysis."],
        ["Production Accounting Engineer", "Maintains monthly hydrocarbon allocation and gas balance."],
    ], input_cols={1}, spans=[1, 7])
    s.gap()
    s.section("(b) Written procedures")
    s.table(["Procedure", "Reference", "Description", "Responsible", "Records location"], [
        ["GHG data collection", f"{f.code}-HSE-PR-110", "Monthly extraction of meter totals from PI historian into GHG workbook.", "GHG Lead", "HSE SharePoint /MRV/2025"],
        ["Meter calibration", f"{f.code}-INS-PR-204", "Annual calibration of emissions meters by accredited lab; out-of-tolerance handling.", "Instr. Supervisor", "CMMS work orders"],
        ["Data gaps and substitution", f"{f.code}-HSE-PR-112", "Substitution using 30-day average; notification to EAD within 30 days; record in H1.", "GHG Lead", "HSE SharePoint /MRV/2025"],
        ["Gas sampling and analysis", f"{f.code}-LAB-PR-031", "Quarterly sampling at SP-3001 and SP-5101; GC per ISO 6974, properties per ISO 6976.", "Lab Coordinator", "LIMS"],
    ], input_cols={2}, spans=[1, 1, 3, 1, 2])
    s.gap()
    s.section("(c) Measuring instruments register")
    rows = []
    for inst in instruments(f):
        rows.append([inst["tag"], inst["service"], inst["type"], inst["tier"], inst["last"].strftime("%d-%b-%Y"),
                     inst["due"].strftime("%d-%b-%Y"), inst["cert"], inst["register_status"]])
    s.table(["Tag", "Service", "Type", "Tier", "Last calibration", "Next due", "Certificate", "Status"], rows, input_cols=set(range(8)))

    s = Sheet(wb, "J_Mitigation_Measures", [8, 40, 16, 16, 10, 16, 40])
    s.title("J. Mitigation measures", f)
    s.section("(a) Actions, plans and studies to reduce GHG emissions")
    s.table(["ID", "Measure", "Type", "Status", "Year", "Est. reduction (t CO2e/yr)", "Notes"], mitigation(f), input_cols=set(range(7)))
    s.gap()
    s.section("(b) Offsets and carbon removals")
    s.para("None used. Offsets and removals are reported for information only and are not netted against reported emissions.")

    s = Sheet(wb, "K_Reference_Lists", [34, 22, 18, 50])
    s.title("K. Reference lists", f)
    s.section("Default factors used")
    s.table(["Parameter", "Value", "Unit", "Source"], [
        ["Natural gas CO2 EF (IPCC default)", IPCC_NG_EF, "t CO2/TJ", "IPCC 2006 Guidelines, Vol. 2, Ch. 2"],
        ["Gas/diesel oil NCV", DIESEL_NCV, "GJ/t", "IPCC 2006 Guidelines, Vol. 2, Ch. 1"],
        ["Gas/diesel oil CO2 EF", DIESEL_EF, "t CO2/TJ", "IPCC 2006 Guidelines, Vol. 2, Ch. 2"],
        ["CH4 EF natural gas combustion", CH4_EF_GAS_COMB, "kg CH4/TJ", "IPCC 2006 Guidelines, Vol. 2, Table 2.2"],
        ["CH4 EF diesel combustion", CH4_EF_DIESEL_COMB, "kg CH4/TJ", "IPCC 2006 Guidelines, Vol. 2, Table 2.2"],
        ["Flare combustion efficiency", FLARE_CE, "fraction", "Industry default for well-operated flares"],
        ["GWP CH4", GWP_CH4, "", "IPCC AR5, 100-year"],
        ["Molar volume", 23.645, "m3/kmol", "Ideal gas at 15 degC, 101.325 kPa"],
        ["Gas to boe", round(SM3_PER_BOE, 1), "Sm3/boe", "5,800 scf per boe"],
    ], fmts={1: "0.000"}, input_cols=set())
    s.gap()
    s.section("Drop-down lists")
    s.table(["List", "Values"], [
        ["Source stream category", "Major; Minor; De-minimis"],
        ["Tier", "Tier 1 (+/-7.5%); Tier 2 (+/-5%); Tier 3 (+/-2.5%); Tier 4 (+/-1.5%); No tier (fallback)"],
        ["Product category", "Refinery products; Coke; Hot metal; Primary aluminium; Grey cement clinker; Lime; Ammonia; Hydrogen; Synthesis gas; "
                             "Aromatics; Heat benchmark; Fuel benchmark; Process emissions; Other"],
    ], input_cols=set(), spans=[1, 3])

    wb.active = 0
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


STREAM_META = {
    "SS-01": {"type": "Combustion", "sources": "ES-01 GTG-A/B/C; ES-02 K-201A/B GT drivers; ES-03 H-101A/B crude heaters",
              "instrument": "FT-3001 ultrasonic meter, fuel gas header", "tier": "Tier 3", "unc": "+/-2.5%",
              "factors": "NCV and EF from quarterly GC analysis (site-specific, Tier 3)"},
    "SS-02": {"type": "Flaring", "sources": "ES-04 HP flare FL-501", "instrument": "FT-5101 ultrasonic flare gas meter",
              "tier": "Tier 2", "unc": "+/-5.0%", "factors": "EF from quarterly GC of flare header gas (Tier 3)"},
    "SS-03": {"type": "Flaring", "sources": "ES-05 LP flare FL-502", "instrument": "FT-5102 ultrasonic flare gas meter",
              "tier": "Tier 2", "unc": "+/-5.0%", "factors": "EF from quarterly GC of flare header gas (Tier 3)"},
    "SS-04": {"type": "Combustion", "sources": "ES-06 EDG-1; ES-07 fire water pumps P-801A/B",
              "instrument": "Supplier invoices and tank dips", "tier": "Tier 1", "unc": "+/-7.5%", "factors": "IPCC 2006 defaults"},
}


def facility_summary(f: Facility) -> str:
    tank_text = ("Crude is stabilised and stored in two fixed-roof tanks (T-401A/B) fitted with pressure/vacuum vents before export by pipeline."
                 if f.flagged else
                 "Crude is stabilised and stored in two fixed-roof tanks (T-401A/B) whose vapours are recovered by VRU-401 before export by pipeline.")
    return (f"{f.name} is an onshore oil processing facility in the {f.field_name}, Al Dhafra Region, operated by {OPERATOR['name']} since "
            f"{f.start_year}. Well fluids are separated in HP and LP inlet separators. {tank_text} Associated gas is compressed by two "
            "gas-turbine-driven compressors (K-201A/B), dehydrated in a TEG unit (U-250) and exported as sales gas or reinjected. "
            "Power is generated on site by three gas turbine generators. Fuel gas is taken from the dehydrated gas stream. "
            "An HP and an LP flare handle relief, purge and upset gas. Diesel is used only for the emergency generator and fire water pumps. "
            "The facility boundary covers all equipment within the CPF fence; wells and flowlines are included, the export pipeline beyond "
            "the custody transfer meter is excluded (operated by a third party).")


def technical_units(f: Facility) -> list[list[str]]:
    rows = [
        ["GTG-A/B/C", "Gas turbine generators", "3 x 24 MW", "SS-01 Fuel gas", "Yes", ""],
        ["K-201A/B", "Gas compressors, gas-turbine driven", "2 x 11 MW", "SS-01 Fuel gas", "Yes", "Wet seals"],
        ["H-101A/B", "Crude heaters (direct fired)", "2 x 18 MW(th)", "SS-01 Fuel gas", "Yes", ""],
        ["FL-501", "HP flare", "Design 4.2 MMSm3/d", "SS-02 HP flare gas", "Yes", "Continuous pilot"],
        ["FL-502", "LP flare", "Design 0.6 MMSm3/d", "SS-03 LP flare gas", "Yes", "Continuous pilot"],
        ["EDG-1 / P-801A/B", "Emergency diesel generator, fire water pumps", "2.5 MW / 2 x 0.8 MW", "SS-04 Diesel", "Yes", "Standby"],
        ["U-250", "TEG dehydration unit", "3.0 MMSm3/d", "-", "Yes", "Still vent (methane source)"],
        ["T-401A/B", "Crude storage tanks", "2 x 250,000 bbl", "-", "Yes", "Tank vents" if f.flagged else "Vapours to VRU-401"],
    ]
    if not f.flagged:
        rows.append(["VRU-401", "Vapour recovery unit (electric)", "0.25 MMSm3/d", "-", "Yes", "Installed 2023"])
    return rows


def instruments(f: Facility) -> list[dict]:
    if f.flagged:
        return [
            {"tag": "FT-3001", "service": "Fuel gas header", "type": "Ultrasonic, 4-path, DN300", "tier": "Tier 3", "range": "0-18,000 Sm3/h",
             "last": date(2025, 1, 18), "due": date(2026, 1, 31), "cert": "GCS-CAL-25-0117", "register_status": "In service",
             "as_found": 0.41, "as_left": 0.12, "limit": 1.0, "serial": "US4-300-18842"},
            {"tag": "FT-5101", "service": "HP flare header", "type": "Ultrasonic flare gas meter, 2-path, DN600", "tier": "Tier 2", "range": "0.03-100 m/s",
             "last": date(2024, 5, 28), "due": date(2025, 5, 31), "cert": "GCS-CAL-24-0562", "register_status": "In service",
             "as_found": 1.9, "as_left": 0.6, "limit": 3.0, "serial": "UFG-600-07731"},
            {"tag": "FT-5102", "service": "LP flare header", "type": "Ultrasonic flare gas meter, 2-path, DN300", "tier": "Tier 2", "range": "0.03-100 m/s",
             "last": date(2025, 2, 10), "due": date(2026, 2, 28), "cert": "GCS-CAL-25-0188", "register_status": "In service",
             "as_found": 1.2, "as_left": 0.5, "limit": 3.0, "serial": "UFG-300-09215"},
        ]
    return [
        {"tag": "FT-3001", "service": "Fuel gas header", "type": "Ultrasonic, 4-path, DN300", "tier": "Tier 3", "range": "0-15,000 Sm3/h",
         "last": date(2025, 1, 22), "due": date(2026, 1, 31), "cert": "GCS-CAL-25-0131", "register_status": "In service",
         "as_found": 0.36, "as_left": 0.10, "limit": 1.0, "serial": "US4-300-17205"},
        {"tag": "FT-5101", "service": "HP flare header", "type": "Ultrasonic flare gas meter, 2-path, DN600", "tier": "Tier 2", "range": "0.03-100 m/s",
         "last": date(2025, 2, 4), "due": date(2026, 2, 28), "cert": "GCS-CAL-25-0152", "register_status": "In service",
         "as_found": 1.4, "as_left": 0.4, "limit": 3.0, "serial": "UFG-600-06984"},
        {"tag": "FT-5102", "service": "LP flare header", "type": "Ultrasonic flare gas meter, 2-path, DN300", "tier": "Tier 2", "range": "0.03-100 m/s",
         "last": date(2025, 3, 14), "due": date(2026, 3, 31), "cert": "GCS-CAL-25-0276", "register_status": "In service (transducer replaced 14-Mar-2025)",
         "as_found": 2.1, "as_left": 0.3, "limit": 3.0, "serial": "UFG-300-08870"},
    ]


def mitigation(f: Facility) -> list[list]:
    if f.flagged:
        return [
            ["MM-01", "GTG-A hot gas path upgrade (efficiency improvement)", "Energy efficiency", "Implemented", 2025, 2_100, "Commissioned November 2025."],
            ["MM-02", "Flare gas recovery unit (FGRU) on HP flare header", "Flaring reduction", "Planned (FEED)", 2026, 45_000, "FEED study to start Q2 2026."],
            ["MM-03", "Leak detection and repair (LDAR) programme", "Methane", "Planned", 2026, "TBD", "Scope under definition."],
        ]
    return [
        ["MM-01", "Vapour recovery unit VRU-401 on crude storage tanks", "Methane", "Implemented", 2023, 80_000, "Recovers tank flash gas to LP compression."],
        ["MM-02", "Quarterly OGI leak detection and repair (LDAR)", "Methane", "Implemented", 2024, 3_500, "Repairs within 15 days of detection."],
        ["MM-03", "Conversion of 118 pneumatic controllers to instrument air", "Methane", "Implemented", 2024, 5_900, "42 low/intermittent-bleed devices remain on gas."],
        ["MM-04", "Flare gas recovery unit (FGRU) on HP flare header", "Flaring reduction", "Planned (EPC award)", 2027, 38_000, "Final investment decision taken Dec 2025."],
    ]
